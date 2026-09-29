"""Plan C: export_all with all four formats (design 11, file names of design 10.6)."""
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from docx import Document as WordDocument
from PIL import Image

from owlocr import export
from owlocr.export import FORMATS, export_all
from owlocr.pipeline.document import Block, Document, Flag, Page
from tests.conftest import PAGES

TEXT = "Rostlina je bylinny strom a także list."


def _doc(source: Path) -> Document:
    flags = [Flag("suspicious", 12, 19, "bylinny", "not in the dictionary"),
             Flag("foreign_letter", 28, 33, "także", "letter does not exist in Czech")]
    blocks = [Block("text", (100, 100, 900, 200), TEXT, TEXT, flags),
              Block("footer", (100, 950, 900, 980), "Zápatí stránky", "Zápatí stránky", [])]
    page = Page(0, "ocr", "quality", 2480, 3508, 0, blocks, "", [], 1.0)
    return Document(str(source), "unlimited_ocr", "x", "0.1.0", "2026-09-28T10:00:00", [page])


def _png(tmp_path: Path) -> Path:
    source = tmp_path / "dopis.png"
    with Image.open(PAGES / "01_letter_clean.png") as img:
        img.save(source, dpi=(200, 200))
    return source


def test_all_four_formats_are_available():
    assert export.FORMATS_AVAILABLE == ("md", "txt", "docx", "pdf") == FORMATS


def test_export_all_writes_all_four_formats_and_no_sidecar(tmp_path):
    source = _png(tmp_path)
    outputs = export_all(_doc(source), source, list(FORMATS), tmp_path / "out", {})
    assert outputs == {"md": tmp_path / "out" / "dopis.md", "txt": tmp_path / "out" / "dopis.txt",
                       "docx": tmp_path / "out" / "dopis.docx",
                       "pdf": tmp_path / "out" / "dopis.ocr.pdf"}
    assert all(p.is_file() for p in outputs.values())
    assert any(p.text == TEXT for p in WordDocument(str(outputs["docx"])).paragraphs)
    pdf = pdfium.PdfDocument(str(outputs["pdf"]))
    assert "Rostlina" in pdf[0].get_textpage().get_text_range()
    pdf.close()
    assert not list((tmp_path / "out").glob("*.owl.json"))      # the runner writes the sidecar


def test_export_all_honours_the_settings(tmp_path):
    source = _png(tmp_path)
    settings = {"keep_page_furniture": True, "append_suspicious_list": True}
    outputs = export_all(_doc(source), source, ["md", "docx"], tmp_path / "out", settings)
    markdown = outputs["md"].read_text(encoding="utf-8")
    assert "Zápatí stránky" in markdown and "- page 1: `także`" in markdown
    assert any(p.text == "Zápatí stránky" for p in WordDocument(str(outputs["docx"])).paragraphs)
    plain = export_all(_doc(source), source, ["docx"], tmp_path / "plain", {})
    assert not any(p.text == "Zápatí stránky"
                   for p in WordDocument(str(plain["docx"])).paragraphs)


def test_exports_use_the_edited_text_not_the_raw_reading(tmp_path):
    source = _png(tmp_path)
    doc = _doc(source)
    doc.pages[0].blocks = [Block("text", (100, 100, 900, 200), "Opravený řádek textu.",
                                 "Puvodni cteni radku.", [])]
    outputs = export_all(doc, source, list(FORMATS), tmp_path / "out", {})
    pdf = pdfium.PdfDocument(str(outputs["pdf"]))
    texts = {
        "md": outputs["md"].read_text(encoding="utf-8"),
        "txt": outputs["txt"].read_text(encoding="utf-8"),
        "docx": "\n".join(p.text for p in WordDocument(str(outputs["docx"])).paragraphs),
        "pdf": pdf[0].get_textpage().get_text_range(),
    }
    pdf.close()
    for fmt, text in texts.items():
        assert "Opravený" in text and "Puvodni" not in text, fmt


def test_unknown_format_is_rejected_before_writing(tmp_path):
    source = _png(tmp_path)
    with pytest.raises(ValueError):
        export_all(_doc(source), source, ["md", "odt"], tmp_path / "out", {})
    assert not (tmp_path / "out" / "dopis.md").exists()
