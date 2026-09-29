"""Plan C: Word export (design 11)."""
from pathlib import Path

from docx import Document as WordDocument
from PIL import Image, ImageDraw

from owlocr.export.docx import export_docx
from owlocr.pipeline.document import Block, Document, Page

TABLE = ("<table><tr><td>Jméno</td><td>Obec</td></tr>"
         "<tr><td>Řehoř Dvořák</td><td>Třebíč</td></tr></table>")
SPAN_TABLE = ("<table><tr><td colspan=\"2\">Přehled</td></tr>"
              "<tr><td>Dub</td><td>Buk</td></tr></table>")
BODY = ("Rostlina je buď velký strom, nebo malý stonek. Listy rostou na stonku a jsou "
        "střídavé nebo vstřícné, podle druhu.")


def _b(label, text, box=(100, 100, 900, 200)):
    return Block(label=label, box=box, text=text, raw_text=text, flags=[])


def _page(index, blocks, width=1000, height=1400, rotation=0):
    return Page(index=index, source="ocr", mode="quality", width_px=width, height_px=height,
                rotation_applied=rotation, blocks=blocks, raw="", warnings=[], seconds=1.0)


def _doc(source: Path, pages):
    return Document(source_path=str(source), engine_id="unlimited_ocr", engine_revision="x",
                    app_version="0.1.0", created="2026-09-28T10:00:00", pages=pages)


def _source_image(path: Path) -> Path:
    img = Image.new("RGB", (1000, 1400), "white")
    ImageDraw.Draw(img).rectangle((200, 700, 600, 1000), fill=(200, 30, 30))   # the "figure"
    img.save(path)
    return path


def _blocks():
    return [
        _b("header", "Učebnice botaniky", (100, 10, 900, 30)),
        _b("title", "Stavba listu", (100, 30, 600, 90)),
        _b("title", "Postavení listů", (100, 100, 500, 120)),
        _b("text", BODY, (100, 130, 900, 160)),
        _b("list", "- první bod\n- druhý bod\n1. očíslovaný bod", (100, 210, 900, 280)),
        _b("table", TABLE, (100, 290, 900, 400)),
        _b("image", "", (200, 500, 600, 714)),
        _b("image_caption", "Obr. 3 Řez listem", (200, 720, 600, 740)),
        _b("formula", "E = mc^2", (100, 750, 400, 780)),
        _b("page_number", "12", (480, 960, 520, 980)),
    ]


def _paragraphs(path):
    return [(p.style.name, p.text) for p in WordDocument(str(path)).paragraphs if p.text]


def test_structure_of_the_word_file(tmp_path):
    source = _source_image(tmp_path / "scan.png")
    out = export_docx(_doc(source, [_page(0, _blocks())]), tmp_path / "scan.docx")
    assert out == tmp_path / "scan.docx"
    paragraphs = _paragraphs(out)
    assert paragraphs[0] == ("Heading 1", "Stavba listu")
    assert paragraphs[1] == ("Heading 2", "Postavení listů")
    assert paragraphs[2] == ("Normal", BODY)
    assert paragraphs[3:6] == [("List Bullet", "první bod"), ("List Bullet", "druhý bod"),
                               ("List Number", "očíslovaný bod")]
    assert ("Normal", "E = mc^2") in paragraphs
    caption = [p for p in WordDocument(str(out)).paragraphs if p.text == "Obr. 3 Řez listem"][0]
    assert all(run.italic for run in caption.runs)


def test_table_is_a_real_table(tmp_path):
    source = _source_image(tmp_path / "scan.png")
    out = export_docx(_doc(source, [_page(0, _blocks())]), tmp_path / "scan.docx")
    table = WordDocument(str(out)).tables[0]
    assert [[c.text for c in row.cells] for row in table.rows] == [
        ["Jméno", "Obec"], ["Řehoř Dvořák", "Třebíč"]]


def test_table_with_spans_still_becomes_a_table(tmp_path):
    source = _source_image(tmp_path / "scan.png")
    doc = _doc(source, [_page(0, [_b("table", SPAN_TABLE)])])
    table = WordDocument(str(export_docx(doc, tmp_path / "s.docx"))).tables[0]
    assert table.cell(1, 0).text == "Dub" and table.cell(1, 1).text == "Buk"
    assert table.cell(0, 0).text == "Přehled"


def test_table_without_html_keeps_its_text(tmp_path):
    doc = _doc(tmp_path / "gone.png", [_page(0, [_b("table", "Jméno Obec\nŘehoř Třebíč")])])
    word = WordDocument(str(export_docx(doc, tmp_path / "t.docx")))
    assert "Řehoř Třebíč" in "\n".join(p.text for p in word.paragraphs)


def test_image_is_cut_from_the_page_and_inlined(tmp_path):
    source = _source_image(tmp_path / "scan.png")
    out = export_docx(_doc(source, [_page(0, _blocks())]), tmp_path / "scan.docx")
    word = WordDocument(str(out))
    assert len(word.inline_shapes) == 1
    blob = word.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed
    image_part = word.part.related_parts[blob]
    from io import BytesIO
    with Image.open(BytesIO(image_part.blob)) as picture:
        assert picture.getpixel((picture.width // 2, picture.height // 2)) == (200, 30, 30)


def test_missing_source_skips_images_but_keeps_text(tmp_path):
    out = export_docx(_doc(tmp_path / "gone.png", [_page(0, _blocks())]), tmp_path / "x.docx")
    word = WordDocument(str(out))
    assert len(word.inline_shapes) == 0
    assert any(p.text == BODY for p in word.paragraphs)


def test_furniture_left_out_unless_asked(tmp_path):
    source = _source_image(tmp_path / "scan.png")
    doc = _doc(source, [_page(0, _blocks())])
    without = [t for _s, t in _paragraphs(export_docx(doc, tmp_path / "a.docx"))]
    assert "Učebnice botaniky" not in without and "12" not in without
    kept = [t for _s, t in _paragraphs(export_docx(doc, tmp_path / "b.docx", keep_furniture=True))]
    assert "Učebnice botaniky" in kept and "12" in kept


def test_optional_page_breaks(tmp_path):
    source = _source_image(tmp_path / "scan.png")
    pages = [_page(0, [_b("text", "První strana.")]), _page(1, [_b("text", "Druhá strana.")])]

    def breaks(path):
        return WordDocument(str(path)).element.body.xml.count('w:type="page"')

    assert breaks(export_docx(_doc(source, pages), tmp_path / "a.docx")) == 0
    assert breaks(export_docx(_doc(source, pages), tmp_path / "b.docx", page_breaks=True)) == 1


def test_existing_file_is_not_overwritten(tmp_path):
    source = _source_image(tmp_path / "scan.png")
    (tmp_path / "scan.docx").write_bytes(b"keep")
    out = export_docx(_doc(source, [_page(0, _blocks())]), tmp_path / "scan.docx")
    assert out == tmp_path / "scan_1.docx"


def test_image_is_cut_from_the_turned_page(tmp_path):
    upright = Image.new("RGB", (1000, 1400), "white")
    ImageDraw.Draw(upright).rectangle((200, 700, 600, 1000), fill=(200, 30, 30))
    source = tmp_path / "sideways.png"
    upright.rotate(90, expand=True).save(source)       # scanned sideways (counter-clockwise)
    page = _page(0, [_b("image", "", (200, 500, 600, 714))], rotation=90)   # read after 90° cw
    word = WordDocument(str(export_docx(_doc(source, [page]), tmp_path / "t.docx")))
    blob = word.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed
    from io import BytesIO
    with Image.open(BytesIO(word.part.related_parts[blob].blob)) as picture:
        assert picture.getpixel((picture.width // 2, picture.height // 2)) == (200, 30, 30)
        assert picture.width > picture.height                 # upright crop, not sideways


def test_xml_invalid_characters_do_not_break_the_word_file(tmp_path):
    # final review C1: python-docx refuses U+FFFE and C0 control characters
    bad = "Rostlina\ufffe je\x0b strom\uffff\x01."
    table = f"<table><tr><td>Dub{chr(0x1f)}</td><td>Buk\ufffe</td></tr></table>"
    blocks = [_b("title", "Stavba\x0c listu"), _b("text", bad), _b("table", table),
              _b("image_caption", "Obr.\x02 1"), _b("list", "- první\x03 bod")]
    out = export_docx(_doc(tmp_path / "missing.png", [_page(0, blocks)]), tmp_path / "x.docx")
    word = WordDocument(str(out))
    texts = [p.text for p in word.paragraphs if p.text]
    assert texts == ["Stavba listu", "Rostlina je strom.", "Obr. 1", "první bod"]
    assert [c.text for c in word.tables[0].rows[0].cells] == ["Dub", "Buk"]


def test_failed_word_write_leaves_no_file(tmp_path, monkeypatch):
    # final review M5: the .docx is written to a temporary name and moved into place
    import docx.document
    import pytest

    def failing_save(self, path):
        Path(path).write_bytes(b"PK half written")
        raise OSError("disk full")

    monkeypatch.setattr(docx.document.Document, "save", failing_save)
    with pytest.raises(OSError):
        export_docx(_doc(tmp_path / "missing.png", [_page(0, [_b("text", BODY)])]),
                    tmp_path / "out" / "x.docx")
    assert list((tmp_path / "out").iterdir()) == []


def test_a_page_that_cannot_be_rendered_skips_pictures_with_a_warning(tmp_path, monkeypatch,
                                                                       caplog):
    # final review T17: only file/image errors are swallowed, and they are logged
    import logging
    import pytest
    from owlocr.export import docx as docx_export

    def broken(*args, **kwargs):
        raise OSError("damaged page")

    monkeypatch.setattr(docx_export, "render_page", broken)
    source = _source_image(tmp_path / "scan.png")
    blocks = [_b("image", "", (200, 500, 600, 714)), _b("text", BODY)]
    with caplog.at_level(logging.WARNING, logger="owlocr.export.docx"):
        out = export_docx(_doc(source, [_page(0, blocks)]), tmp_path / "scan.docx")
    assert len(WordDocument(str(out)).inline_shapes) == 0
    assert "damaged page" in caplog.text

    def bug(*args, **kwargs):
        raise RuntimeError("a programming error")

    monkeypatch.setattr(docx_export, "render_page", bug)
    with pytest.raises(RuntimeError):
        export_docx(_doc(source, [_page(0, blocks)]), tmp_path / "again.docx")
