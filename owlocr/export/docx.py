"""Word export (design 11): headings, paragraphs, lists, real tables, images cut from the page
image, captions in italics, optional page breaks. python-docx (MIT)."""
from __future__ import annotations

import io
import logging
import re
import statistics
import tempfile
from pathlib import Path

from docx import Document as WordDocument
from docx.enum.text import WD_BREAK
from docx.shared import Inches
from PIL import Image

from owlocr.paths import atomic_write_file, unique_path
from owlocr.pipeline.document import FURNITURE_LABELS, Block, Document, Page
from owlocr.pipeline.guards import rotate_image
from owlocr.pipeline.pages import render_page, xml_safe
from owlocr.pipeline.parse import html_table_to_rows

log = logging.getLogger(__name__)

IMAGE_DPI = 150
TEXT_WIDTH_IN = 6.0            # usable width of the default Word page
HEADING_1_RATIO = 1.6          # same rule as the Markdown export of plan A
_BULLET = re.compile(r"^\s*(?:[-*•·▪◦]|(\d+|[a-z])[.)])\s+")
_ROW = re.compile(r"<tr\b[^>]*>(.*?)</tr>", re.I | re.S)
_CELL = re.compile(r"<t[dh]\b[^>]*>(.*?)</t[dh]>", re.I | re.S)
_TAG = re.compile(r"<[^>]*>")
_CAPTIONS = ("image_caption", "table_caption")
_PICTURES = ("image", "figure")


def _line_height(page: Page) -> float | None:
    """Typical height of one text line (0..999 units), the rule of plan A's Markdown export:
    the median of the smaller half of the text blocks' box heights."""
    heights = sorted(b.box[3] - b.box[1] for b in page.blocks if b.label == "text" and b.box)
    if not heights:
        return None
    return statistics.median(heights[: max(1, len(heights) // 2)])


def _heading_level(block: Block, line: float | None) -> int:
    """1 when the title's box is at least 1.6 times a text line (as '#' in Markdown), else 2."""
    if block.box and line and (block.box[3] - block.box[1]) >= HEADING_1_RATIO * line:
        return 1
    return 2


def _table_rows(html: str) -> list[list[str]]:
    rows = html_table_to_rows(html)
    if rows is not None:
        return rows
    # rowspan/colspan present: keep the cells in reading order and ignore the spans
    return [[" ".join(_TAG.sub(" ", cell).split()) for cell in _CELL.findall(row)]
            for row in _ROW.findall(html)]


def _add_text(paragraph, text: str, italic: bool = False) -> None:
    text = xml_safe(text)          # python-docx refuses XML-invalid characters (C1)
    for n, line in enumerate(text.split("\n")):
        run = paragraph.add_run(line)
        run.italic = italic or None
        if n < text.count("\n"):
            run.add_break(WD_BREAK.LINE)


def _add_table(word, block: Block) -> None:
    # ruling R4: a `table` block whose text is not HTML keeps its text as a paragraph
    rows = [r for r in _table_rows(block.text) if r] if "<t" in block.text.lower() else []
    if not rows:                       # not HTML (or empty): keep the text as a paragraph
        if block.text.strip():
            _add_text(word.add_paragraph(), block.text.strip())
        return
    columns = max(len(r) for r in rows)
    table = word.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            table.cell(r, c).text = xml_safe(cell)


def _add_list(word, block: Block) -> None:
    for line in block.text.split("\n"):
        if not line.strip():
            continue
        match = _BULLET.match(line)
        style = "List Number" if match and match.group(1) else "List Bullet"
        word.add_paragraph(xml_safe(line[match.end():] if match else line.strip()), style=style)


class _PageImages:
    """Renders each page image once (upright, as the model read it) for cutting out pictures."""

    def __init__(self, source: Path, folder: Path):
        self.source = source
        self.folder = folder
        self.cache: dict[int, Image.Image | None] = {}

    def get(self, page: Page) -> Image.Image | None:
        if page.index not in self.cache:
            self.cache[page.index] = self._render(page)
        return self.cache[page.index]

    def _render(self, page: Page) -> Image.Image | None:
        if not self.source.is_file():
            return None
        png = self.folder / f"page_{page.index:04d}.png"
        try:
            render_page(self.source, page.index, IMAGE_DPI, png)
            if page.rotation_applied:
                png = rotate_image(png, page.rotation_applied, png.with_name(png.stem + "_up.png"))
            with Image.open(png) as img:
                img.load()
                return img.copy()
        except (OSError, ValueError, IndexError, EOFError) as exc:     # final review T17
            log.warning("pictures of page %d left out: the page could not be rendered (%s)",
                        page.index + 1, exc)
            return None


def _add_picture(word, block: Block, page: Page, images: _PageImages) -> None:
    image = images.get(page)
    if image is None or block.box is None:
        return
    x1, y1, x2, y2 = block.box
    crop = image.crop((round(x1 / 999 * image.width), round(y1 / 999 * image.height),
                       round(x2 / 999 * image.width), round(y2 / 999 * image.height)))
    if crop.width < 2 or crop.height < 2:
        return
    stream = io.BytesIO()
    crop.convert("RGB").save(stream, format="PNG")
    stream.seek(0)
    width = max(0.5, min(TEXT_WIDTH_IN, (x2 - x1) / 999 * TEXT_WIDTH_IN))
    word.add_picture(stream, width=Inches(width))


def export_docx(doc: Document, out: Path, page_breaks: bool = False,
                keep_furniture: bool = False) -> Path:
    out = unique_path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    word = WordDocument()
    with tempfile.TemporaryDirectory() as tmp:
        images = _PageImages(Path(doc.source_path), Path(tmp))
        for n, page in enumerate(doc.pages):
            if n and page_breaks:
                word.add_page_break()
            line = _line_height(page)
            for block in page.blocks:
                if block.label in FURNITURE_LABELS and not keep_furniture:
                    continue
                if block.label in _PICTURES:
                    _add_picture(word, block, page, images)
                elif not block.text.strip():
                    continue
                elif block.label == "title":
                    level = _heading_level(block, line)
                    word.add_heading(xml_safe(block.text.replace("\n", " ")), level=level)
                elif block.label == "table":
                    _add_table(word, block)
                elif block.label == "list":
                    _add_list(word, block)
                elif block.label in _CAPTIONS:
                    _add_text(word.add_paragraph(), block.text, italic=True)
                else:
                    _add_text(word.add_paragraph(), block.text)
        atomic_write_file(out, lambda tmp: word.save(str(tmp)))     # final review M5
    return out
